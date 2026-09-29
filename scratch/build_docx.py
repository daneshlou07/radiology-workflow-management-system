import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(val='single', sz='4', space='0', color='CCCCCC')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'D3D3D3'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def build_report():
    doc = Document()

    # Page Margins (Standard UNITEN thesis: 1.5 in Left, 1.0 in Right/Top/Bottom)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    # Read markdown content
    md_path = r'c:\Users\Danesh Lou\Downloads\Rebuild\docs\FYP1_REPORT_DRAFT_HEALTHGRID_IQ.md'
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # parse rows
        parsed_rows = []
        for r in table_rows:
            # split by pipe
            cells = [c.strip() for c in r.strip('|').split('|')]
            # check if separator row
            if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                continue
            parsed_rows.append(cells)
        
        if not parsed_rows:
            in_table = False
            table_rows = []
            return

        num_cols = max(len(r) for r in parsed_rows)
        # Pad shorter rows
        for r in parsed_rows:
            while len(r) < num_cols:
                r.append('')

        tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for row_idx, r_data in enumerate(parsed_rows):
            is_header = (row_idx == 0)
            row = tbl.rows[row_idx]
            for col_idx, cell_value in enumerate(r_data):
                cell = row.cells[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                set_cell_border(cell, 
                                top=dict(val='single', sz=6 if is_header else 4, color='94A3B8' if is_header else 'E2E8F0'),
                                bottom=dict(val='single', sz=8 if is_header else 4, color='475569' if is_header else 'E2E8F0'),
                                left=dict(val='single', sz=4, color='E2E8F0'),
                                right=dict(val='single', sz=4, color='E2E8F0'))
                
                if is_header:
                    set_cell_shading(cell, 'F1F5F9')

                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                
                # strip bold markdown if any
                clean_val = cell_value.replace('**', '')
                run = p.add_run(clean_val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10 if num_cols > 3 else 11)
                if is_header:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                else:
                    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        doc.add_paragraph() # space after table
        in_table = False
        table_rows = []

    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()

        # Check table
        if trimmed.startswith('|') and trimmed.endswith('|'):
            in_table = True
            table_rows.append(trimmed)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Check page break
        if '\\pagebreak' in trimmed or trimmed == '---':
            # Check if this is horizontal rule or pagebreak
            if '\\pagebreak' in trimmed:
                doc.add_page_break()
                i += 1
                continue
            elif i > 0 and lines[i-1].strip().startswith('#'):
                # just a divider under heading
                i += 1
                continue
            elif trimmed == '---':
                # divider
                i += 1
                continue

        # Skip empty lines
        if not trimmed:
            i += 1
            continue

        # Check Headings
        if trimmed.startswith('# '):
            text = trimmed[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.keep_with_next = True
            if 'CHAPTER' in text.upper() or 'UNIVERSITI' in text.upper() or 'HEALTHGRID' in text.upper():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif trimmed.startswith('## '):
            text = trimmed[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            if any(k in text.upper() for k in ['DECLARATION', 'APPROVAL PAGE', 'ACKNOWLEDGMENTS', 'ABSTRACT', 'TABLE OF CONTENTS', 'LIST OF']):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        elif trimmed.startswith('### '):
            text = trimmed[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif trimmed.startswith('#### '):
            text = trimmed[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.italic = True
        elif trimmed.startswith('* ') or trimmed.startswith('- '):
            text = trimmed[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(4)
            # parse bold inside bullet
            add_formatted_text(p, text)
        elif re.match(r'^\d+\.\s+', trimmed):
            match = re.match(r'^(\d+\.)\s+(.*)', trimmed)
            num_prefix = match.group(1)
            text = match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.25)
            r_num = p.add_run(num_prefix + ' ')
            r_num.font.name = 'Times New Roman'
            r_num.font.bold = True
            add_formatted_text(p, text)
        elif trimmed.startswith('>'):
            text = trimmed.lstrip('> ').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.italic = True
        elif trimmed.startswith('$$') and trimmed.endswith('$$'):
            formula = trimmed.strip('$').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(formula)
            run.font.name = 'Cambria Math'
            run.font.size = Pt(11)
            run.font.italic = True
        elif trimmed.startswith('```'):
            # code block or text diagram
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        else:
            p = doc.add_paragraph()
            # Check for title page metadata centering
            if any(k in trimmed for k in ['MUHAMMAD DANESH HAKIMI LOU', 'BSW01084693', 'Dr. Siti Rohana', 'College of Computing', 'Bachelor of Computer Science', 'SEPTEMBER 2026', 'By:']):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.3
            add_formatted_text(p, trimmed)

        i += 1

    if in_table:
        flush_table()

    # Save outputs
    out_paths = [
        r'c:\Users\Danesh Lou\Downloads\HealthGrid_IQ_FYP1_Report.docx',
        r'c:\Users\Danesh Lou\Downloads\FYP 1\HealthGrid_IQ_FYP1_Report.docx'
    ]
    for op in out_paths:
        try:
            os.makedirs(os.path.dirname(op), exist_ok=True)
            doc.save(op)
            print(f'Successfully saved: {op}')
        except Exception as e:
            print(f'Error saving {op}: {e}')

def add_formatted_text(paragraph, text):
    # Regex split for **bold** and *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = 'Times New Roman'
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Times New Roman'
            run.font.italic = True
        else:
            run = paragraph.add_run(token)
            run.font.name = 'Times New Roman'

if __name__ == '__main__':
    build_report()
